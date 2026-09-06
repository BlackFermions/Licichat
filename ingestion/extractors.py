"""Bounded document extraction with selective OCR and safe archive handling."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path, PurePosixPath

try:
    import fitz
except ImportError:  # Pure archive tests do not require the PDF runtime.
    fitz = None

try:
    import pytesseract
    from PIL import Image, ImageOps
except ImportError:  # Reported explicitly if an OCR page is encountered.
    pytesseract = None
    Image = None
    ImageOps = None

try:
    from docx import Document
except ImportError:
    Document = None

from ingestion.config import Settings
from ingestion.errors import IngestionError
from ingestion.models import ExtractedPage, ExtractionResult


SUPPORTED_MEMBER_SUFFIXES = {".pdf", ".docx", ".txt"}
FORMAT_SUFFIXES = {
    "application/pdf": ".pdf",
    "pdf": ".pdf",
    "application/zip": ".zip",
    "zip": ".zip",
    "application/x-rar-compressed": ".rar",
    "rar": ".rar",
    "7z": ".7z",
    "application/x-7z-compressed": ".7z",
    "docx": ".docx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
}


def source_suffix(source_format: str | None, title: str | None) -> str:
    normalized = str(source_format or "").strip().lower()
    if normalized in FORMAT_SUFFIXES:
        return FORMAT_SUFFIXES[normalized]
    title_suffix = Path(str(title or "")).suffix.lower()
    if title_suffix in {".pdf", ".zip", ".rar", ".7z", ".docx", ".txt"}:
        return title_suffix
    return ".pdf"


def detect_kind(path: str) -> str:
    file_path = Path(path)
    with file_path.open("rb") as handle:
        prefix = handle.read(8)
    if prefix.startswith(b"%PDF-"):
        return "pdf"
    if prefix.startswith(b"Rar!\x1a\x07"):
        return "rar"
    if prefix.startswith(b"7z\xbc\xaf\x27\x1c"):
        return "7z"
    if prefix.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        try:
            with zipfile.ZipFile(file_path) as archive:
                names = {name.replace("\\", "/") for name in archive.namelist()}
            if "[Content_Types].xml" in names and "word/document.xml" in names:
                return "docx"
        except zipfile.BadZipFile as exc:
            raise IngestionError("invalid_zip", "El archivo ZIP no es valido.") from exc
        return "zip"
    if file_path.suffix.lower() == ".txt":
        return "txt"
    raise IngestionError("unsupported_format", "El archivo no tiene un formato compatible.")


def _safe_member_name(raw_name: str) -> str:
    normalized = str(raw_name or "").replace("\\", "/").strip()
    if not normalized or "\x00" in normalized:
        raise IngestionError("unsafe_archive", "El archivo comprimido contiene una ruta invalida.")
    path = PurePosixPath(normalized)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise IngestionError("unsafe_archive", "El archivo comprimido contiene una ruta insegura.")
    return normalized


def _member_rank(name: str, role: str) -> tuple[int, int, str]:
    normalized = name.lower()
    if role == "buena_pro":
        preferred = ("buena pro", "otorgamiento", "acta")
    else:
        preferred = ("bases integradas", "bases administrativas", "bases", "estandar")
    score = next((index for index, token in enumerate(preferred) if token in normalized), 20)
    suffix_rank = {".pdf": 0, ".docx": 1, ".txt": 2}.get(Path(name).suffix.lower(), 9)
    return score, suffix_rank, normalized


def _select_members(names: list[str], role: str, settings: Settings) -> list[str]:
    if len(names) > settings.archive_max_files:
        raise IngestionError("archive_too_many_files", "El archivo comprimido contiene demasiados elementos.")
    candidates = []
    for raw_name in names:
        name = _safe_member_name(raw_name)
        if Path(name).suffix.lower() in SUPPORTED_MEMBER_SUFFIXES:
            candidates.append(name)
    if not candidates:
        raise IngestionError("archive_no_supported_files", "El archivo no contiene PDF, DOCX o TXT compatibles.")
    candidates.sort(key=lambda name: _member_rank(name, role))
    preferred = [name for name in candidates if _member_rank(name, role)[0] < 20]
    selected = preferred or candidates
    return selected[: settings.archive_max_selected_files]


def _read_zip_members(path: str, role: str, settings: Settings) -> list[tuple[str, bytes]]:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = [info for info in archive.infolist() if not info.is_dir()]
            names = _select_members([info.filename for info in infos], role, settings)
            by_name = {_safe_member_name(info.filename): info for info in infos}
            selected_infos = [by_name[name] for name in names]
            total_size = sum(info.file_size for info in selected_infos)
            if total_size > settings.archive_max_uncompressed_bytes:
                raise IngestionError("archive_too_large", "El contenido descomprimido supera el limite.")
            for info in selected_infos:
                ratio = info.file_size / max(info.compress_size, 1)
                if ratio > settings.archive_max_ratio:
                    raise IngestionError("archive_ratio_exceeded", "El archivo comprimido excede el ratio permitido.")
            return [(name, archive.read(by_name[name])) for name in names]
    except zipfile.BadZipFile as exc:
        raise IngestionError("invalid_zip", "El archivo ZIP no es valido.") from exc


def _is_regular_archive_entry(entry: object) -> bool:
    marker = getattr(entry, "isfile", False)
    return bool(marker() if callable(marker) else marker)


def _read_libarchive_members(path: str, role: str, settings: Settings) -> list[tuple[str, bytes]]:
    try:
        import libarchive
    except (ImportError, OSError) as exc:
        raise IngestionError("archive_runtime_missing", "El lector RAR/7z no esta disponible.") from exc

    metadata: list[tuple[str, int]] = []
    try:
        with libarchive.file_reader(path) as archive:
            for entry in archive:
                if not _is_regular_archive_entry(entry):
                    continue
                name = _safe_member_name(entry.pathname)
                metadata.append((name, max(0, int(entry.size or 0))))
                if len(metadata) > settings.archive_max_files:
                    raise IngestionError(
                        "archive_too_many_files", "El archivo comprimido contiene demasiados elementos."
                    )
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("invalid_archive", "No se pudo inspeccionar el archivo comprimido.") from exc

    selected_names = _select_members([name for name, _ in metadata], role, settings)
    size_by_name = dict(metadata)
    total_size = sum(size_by_name.get(name, 0) for name in selected_names)
    source_size = max(Path(path).stat().st_size, 1)
    if total_size > settings.archive_max_uncompressed_bytes:
        raise IngestionError("archive_too_large", "El contenido descomprimido supera el limite.")
    if total_size / source_size > settings.archive_max_ratio:
        raise IngestionError("archive_ratio_exceeded", "El archivo comprimido excede el ratio permitido.")

    selected = set(selected_names)
    payloads: dict[str, bytes] = {}
    try:
        with libarchive.file_reader(path) as archive:
            for entry in archive:
                name = _safe_member_name(entry.pathname)
                if name not in selected:
                    continue
                output = bytearray()
                for block in entry.get_blocks():
                    output.extend(block)
                    if len(output) > settings.archive_max_uncompressed_bytes:
                        raise IngestionError("archive_too_large", "El contenido descomprimido supera el limite.")
                payloads[name] = bytes(output)
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("invalid_archive", "No se pudo leer el archivo comprimido.") from exc
    return [(name, payloads[name]) for name in selected_names if name in payloads]


def _clean_text(value: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    output: list[str] = []
    previous_blank = False
    for line in lines:
        if not line:
            if output and not previous_blank:
                output.append("")
            previous_blank = True
            continue
        output.append(line)
        previous_blank = False
    return "\n".join(output).strip()


def _ocr_page(page: "fitz.Page", dpi: int) -> str:
    if pytesseract is None or Image is None or ImageOps is None:
        raise IngestionError("ocr_runtime_missing", "El motor OCR no esta disponible.")
    pixmap = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY, alpha=False)
    image = Image.frombytes("L", (pixmap.width, pixmap.height), pixmap.samples)
    image = ImageOps.autocontrast(image)
    return _clean_text(pytesseract.image_to_string(image, lang="spa+eng", config="--oem 1"))


def _page_image_coverage(page: "fitz.Page") -> float:
    """Estimate image coverage to catch scanned bodies with selectable headers."""
    try:
        rect = page.rect
        page_area = max(float(rect.width) * float(rect.height), 1.0)
        image_area = 0.0
        for info in page.get_image_info() or []:
            bbox = info.get("bbox") if isinstance(info, dict) else None
            if not bbox or len(bbox) < 4:
                continue
            x0, y0, x1, y1 = (float(value) for value in bbox[:4])
            image_area += max(0.0, x1 - x0) * max(0.0, y1 - y0)
        return min(image_area / page_area, 1.0)
    except (AttributeError, TypeError, ValueError):
        return 0.0


def _body_text_chars(page: "fitz.Page") -> int | None:
    try:
        rect = page.rect
        top = float(rect.y0) + float(rect.height) * 0.12
        bottom = float(rect.y0) + float(rect.height) * 0.88
        total = 0
        for block in page.get_text("blocks", sort=True) or []:
            if not isinstance(block, (tuple, list)) or len(block) < 5:
                continue
            y0, y1 = float(block[1]), float(block[3])
            if y1 <= top or y0 >= bottom:
                continue
            total += len(re.sub(r"\s+", "", str(block[4] or "")))
        return total
    except (AttributeError, TypeError, ValueError):
        return None


def _requires_ocr(page: "fitz.Page", native: str, minimum_chars: int) -> bool:
    native_chars = len(re.sub(r"\s+", "", native))
    if native_chars < minimum_chars:
        return True
    body_chars = _body_text_chars(page)
    if body_chars is None:
        return False
    # Some SEACE PDFs expose headers and thousands of layout spaces as a text
    # layer while the meaningful center of the page remains an image. Body
    # sparsity is therefore a stronger signal than raw image metadata, which
    # varies between PDF producers and PyMuPDF versions.
    return (
        native_chars < max(600, minimum_chars * 6)
        and body_chars < max(120, minimum_chars)
    )


def _extract_pdf(data: bytes, name: str, settings: Settings) -> ExtractionResult:
    if fitz is None:
        raise IngestionError("pdf_runtime_missing", "El lector PDF no esta disponible.")
    pages: list[ExtractedPage] = []
    native_pages = 0
    ocr_pages = 0
    skipped_pages = 0
    warnings: list[str] = []
    total_chars = 0
    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise IngestionError("encrypted_pdf", "El PDF requiere una contrasena.")
            page_count = min(len(document), settings.max_pages_per_document)
            if len(document) > page_count:
                warnings.append("page_limit_reached")
            for index in range(page_count):
                if total_chars >= settings.max_text_chars_per_document:
                    warnings.append("text_limit_reached")
                    break
                page = document[index]
                native = _clean_text(page.get_text("text", sort=True))
                used_ocr = False
                text = native
                if _requires_ocr(page, native, settings.native_text_min_chars):
                    if ocr_pages < settings.ocr_max_pages:
                        ocr_text = _ocr_page(page, settings.ocr_dpi)
                        ocr_pages += 1
                        if len(re.sub(r"\s+", "", ocr_text)) >= settings.native_text_min_chars:
                            text = ocr_text
                            used_ocr = True
                        elif native:
                            if "ocr_low_text" not in warnings:
                                warnings.append("ocr_low_text")
                        else:
                            skipped_pages += 1
                            if "ocr_low_text" not in warnings:
                                warnings.append("ocr_low_text")
                            continue
                    else:
                        skipped_pages += 1
                        if "ocr_limit_reached" not in warnings:
                            warnings.append("ocr_limit_reached")
                        continue
                if not text:
                    skipped_pages += 1
                    continue
                remaining = settings.max_text_chars_per_document - total_chars
                text = text[:remaining]
                pages.append(ExtractedPage(name[:240], index + 1, text, used_ocr))
                total_chars += len(text)
                if not used_ocr:
                    native_pages += 1
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("invalid_pdf", "No se pudo leer el PDF.") from exc
    return ExtractionResult(
        pages=pages,
        detected_mime_type="application/pdf",
        extracted_filename=name,
        page_count=page_count,
        native_text_pages=native_pages,
        ocr_pages=ocr_pages,
        skipped_pages=skipped_pages,
        warnings=warnings,
    )


def _extract_docx(data: bytes, name: str) -> ExtractionResult:
    if Document is None:
        raise IngestionError("docx_runtime_missing", "El lector DOCX no esta disponible.")
    try:
        document = Document(io.BytesIO(data))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = _clean_text(paragraph.text)
            if text:
                blocks.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [_clean_text(cell.text) for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        text = "\n\n".join(blocks)
    except Exception as exc:
        raise IngestionError("invalid_docx", "No se pudo leer el DOCX.") from exc
    pages = [ExtractedPage(name[:240], 1, text)] if text else []
    return ExtractionResult(
        pages=pages,
        detected_mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extracted_filename=name,
        page_count=1,
        native_text_pages=1 if pages else 0,
        ocr_pages=0,
    )


def _extract_text(data: bytes, name: str) -> ExtractionResult:
    text = _clean_text(data.decode("utf-8", errors="replace"))
    pages = [ExtractedPage(name[:240], 1, text)] if text else []
    return ExtractionResult(
        pages=pages,
        detected_mime_type="text/plain",
        extracted_filename=name,
        page_count=1,
        native_text_pages=1 if pages else 0,
        ocr_pages=0,
    )


def _extract_payload(data: bytes, name: str, settings: Settings) -> ExtractionResult:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(data, name, settings)
    if suffix == ".docx":
        return _extract_docx(data, name)
    if suffix == ".txt":
        return _extract_text(data, name)
    raise IngestionError("unsupported_member", "El documento interno no es compatible.")


def _merge_results(results: list[ExtractionResult], source_name: str) -> ExtractionResult:
    if not results:
        raise IngestionError("no_extractable_text", "No se encontro contenido procesable.")
    return ExtractionResult(
        pages=[page for result in results for page in result.pages],
        detected_mime_type="application/x-archive",
        extracted_filename=source_name,
        page_count=sum(result.page_count for result in results),
        native_text_pages=sum(result.native_text_pages for result in results),
        ocr_pages=sum(result.ocr_pages for result in results),
        skipped_pages=sum(result.skipped_pages for result in results),
        warnings=[warning for result in results for warning in result.warnings],
    )


def extract_source(path: str, source_name: str, role: str, settings: Settings) -> ExtractionResult:
    kind = detect_kind(path)
    data = Path(path).read_bytes()
    if kind == "pdf":
        return _extract_pdf(data, source_name, settings)
    if kind == "docx":
        return _extract_docx(data, source_name)
    if kind == "txt":
        return _extract_text(data, source_name)
    if kind == "zip":
        members = _read_zip_members(path, role, settings)
    elif kind in {"rar", "7z"}:
        members = _read_libarchive_members(path, role, settings)
    else:
        raise IngestionError("unsupported_format", "El archivo no tiene un formato compatible.")
    results = [_extract_payload(payload, name, settings) for name, payload in members]
    return _merge_results(results, source_name)
